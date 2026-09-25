from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, Field

from app.api.auth import get_current_user
from app.core.supabase import get_admin_client
from app.services.usage import check_and_increment, try_persist_usage_event

router = APIRouter(prefix="/api/v1/documents", tags=["documents"])


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _http_or_500(exc: Exception, fallback: str) -> HTTPException:
    if isinstance(exc, HTTPException):
        return exc
    return HTTPException(
        status_code=500,
        detail=f"{fallback} ({type(exc).__name__}: {exc})",
    )


class CorrectionRequest(BaseModel):
    """Human correction of AI-produced fields."""
    category: str | None = Field(None, max_length=120)
    summary: str | None = Field(None, max_length=8000)
    structured_data: dict[str, Any] | None = None


@router.get("/export.xlsx")
async def export_documents_xlsx(user=Depends(get_current_user)):
    """Export all of the user's documents as a native Excel workbook."""
    check_and_increment(str(user.id), "exports")
    try_persist_usage_event(str(user.id), "exports")

    try:
        response = (
            get_admin_client()
            .table("documents")
            .select(
                "id,filename,content_type,file_size,status,category,"
                "classification_confidence,summary,error_message,retry_count,"
                "processed_at,structured_data,created_at,updated_at"
            )
            .eq("user_id", str(user.id))
            .order("created_at", desc=True)
            .execute()
        )
    except Exception as exp:
        raise _http_or_500(exp, "Unable to export documents") from exp

    rows = response.data or []
    wb = Workbook()
    ws = wb.active
    ws.title = "Documents"
    headers = [
        "id", "filename", "content_type", "file_size", "status", "category",
        "classification_confidence", "summary", "error_message", "retry_count",
        "processed_at", "structured_data", "created_at", "updated_at",
    ]
    ws.append(headers)
    for row in rows:
        structured = row.get("structured_data")
        ws.append([
            row.get("id"),
            row.get("filename"),
            row.get("content_type"),
            row.get("file_size"),
            row.get("status"),
            row.get("category"),
            row.get("classification_confidence"),
            row.get("summary"),
            row.get("error_message"),
            row.get("retry_count"),
            row.get("processed_at"),
            json.dumps(structured or {}, ensure_ascii=False)
            if structured is not None
            else "",
            row.get("created_at"),
            row.get("updated_at"),
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="documents-export.xlsx"',
        },
    )


@router.get("/{document_id}/export.xlsx")
async def export_document_xlsx(document_id: str, user=Depends(get_current_user)):
    """Export a single document as a native Excel workbook."""
    check_and_increment(str(user.id), "exports")
    try_persist_usage_event(str(user.id), "exports", document_id=document_id)

    try:
        response = (
            get_admin_client()
            .table("documents")
            .select(
                "id,filename,content_type,file_size,status,category,"
                "classification_confidence,summary,error_message,retry_count,"
                "processed_at,structured_data,extracted_text,created_at,updated_at"
            )
            .eq("id", document_id)
            .eq("user_id", str(user.id))
            .maybe_single()
            .execute()
        )
    except Exception as exp:
        raise _http_or_500(exp, "Unable to export document") from exp

    if not response.data:
        raise HTTPException(status_code=404, detail="Document not found.")

    doc = response.data
    wb = Workbook()
    meta = wb.active
    meta.title = "Metadata"
    meta.append(["Field", "Value"])
    for key in (
        "id", "filename", "content_type", "file_size", "status", "category",
        "classification_confidence", "summary", "error_message", "retry_count",
        "processed_at", "created_at", "updated_at",
    ):
        meta.append([key, doc.get(key)])

    fields = wb.create_sheet("Structured fields")
    fields.append(["Key", "Value"])
    structured = doc.get("structured_data") or {}
    if isinstance(structured, dict):
        for k, v in structured.items():
            fields.append([
                k,
                json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v,
            ])

    text_sheet = wb.create_sheet("Extracted text")
    text_sheet.append(["extracted_text"])
    text_sheet.append([doc.get("extracted_text") or ""])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe_name = (doc.get("filename") or "document").rsplit(".", 1)[0]
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{safe_name}-export.xlsx"',
        },
    )


@router.patch("/{document_id}")
async def correct_document(
    document_id: str,
    body: CorrectionRequest,
    user=Depends(get_current_user),
):
    """Human correction: update category, summary, and/or structured_data."""
    updates: dict[str, Any] = {"updated_at": _now_iso()}
    if body.category is not None:
        updates["category"] = body.category.strip() or "other"
    if body.summary is not None:
        updates["summary"] = body.summary.strip()
    if body.structured_data is not None:
        updates["structured_data"] = body.structured_data

    if len(updates) == 1:
        raise HTTPException(status_code=400, detail="No correction fields provided.")

    try:
        response = (
            get_admin_client()
            .table("documents")
            .update(updates)
            .eq("id", document_id)
            .eq("user_id", str(user.id))
            .execute()
        )
    except Exception as exp:
        raise _http_or_500(exp, "Unable to save correction") from exp

    if not response.data:
        check = (
            get_admin_client()
            .table("documents")
            .select("id")
            .eq("id", document_id)
            .eq("user_id", str(user.id))
            .maybe_single()
            .execute()
        )
        if not check.data:
            raise HTTPException(status_code=404, detail="Document not found.")
        refreshed = (
            get_admin_client()
            .table("documents")
            .select(
                "id,filename,content_type,file_size,status,category,"
                "classification_confidence,summary,error_message,retry_count,"
                "processed_at,structured_data,extracted_text,storage_path,"
                "created_at,updated_at"
            )
            .eq("id", document_id)
            .eq("user_id", str(user.id))
            .maybe_single()
            .execute()
        )
        return refreshed.data

    return response.data[0] if isinstance(response.data, list) else response.data
